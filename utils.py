"""Utility functions for the bot"""
import html
import io
import math
from datetime import date, datetime
from typing import Optional, Tuple
from aiogram import types
from logger import logger

MAX_NAME_LENGTH = 100
MAX_NOTES_LENGTH = 500


def esc(value) -> str:
    """Escape user-supplied text for Telegram's HTML parse mode.

    A bare '<' or '&' in a client name makes Telegram reject the whole message.
    """
    return html.escape(str(value), quote=False)


def parse_date(date_str: str) -> Optional[date]:
    """Parse date from string in DD.MM.YYYY format"""
    try:
        return datetime.strptime(date_str.strip(), "%d.%m.%Y").date()
    except ValueError:
        return None


def format_date(date_obj: date) -> str:
    """Format date to DD.MM.YYYY"""
    return date_obj.strftime("%d.%m.%Y")


def format_currency(amount: float, currency: str = "RUB") -> str:
    """Format currency with symbol"""
    symbols = {
        "RUB": "₽",
        "USD": "$",
        "UZS": "сум"
    }
    symbol = symbols.get(currency, currency)
    return f"{amount:,.2f} {symbol}"


def get_days_until(target_date: date) -> int:
    """Get number of days until target date (negative if overdue)"""
    return (target_date - date.today()).days


def get_status_emoji(payment_status: str) -> str:
    """Get emoji for payment status"""
    status_map = {
        'pending': '⏳',
        'paid': '✅',
        'unpaid': '❌',
        'overdue': '⚠️'
    }
    return status_map.get(payment_status, '❓')


def format_client_info(client, show_details: bool = True) -> str:
    """Format client information for display"""
    next_payment = datetime.strptime(client['next_payment_date'], "%Y-%m-%d").date()
    days_until = get_days_until(next_payment)

    # Determine status
    if days_until < 0:
        status = f"⚠️ Просрочено на {abs(days_until)} дн."
        status_emoji = "🔴"
    elif days_until == 0:
        status = "📅 Сегодня"
        status_emoji = "🟡"
    elif days_until <= 3:
        status = f"⏰ Через {days_until} дн."
        status_emoji = "🟡"
    else:
        status = f"✅ Через {days_until} дн."
        status_emoji = "🟢"

    info = (
        f"{status_emoji} <b>{esc(client['name'])}</b>\n"
        f"💰 Сумма: {format_currency(client['amount'], client['currency'])}\n"
        f"📅 Следующий платеж: {format_date(next_payment)}\n"
        f"📊 Статус: {status}\n"
    )

    if show_details:
        info += f"🔄 Интервал: {get_interval_description(client['interval_type'], client['interval_value'])}\n"
        if client['category'] and client['category'] != 'general':
            info += f"🏷️ Категория: {esc(client['category'])}\n"
        if client['notes']:
            info += f"📝 Заметки: {esc(client['notes'])}\n"

    return info


def validate_name(name_str: str) -> Tuple[bool, Optional[str], str]:
    """Validate and normalize a client name"""
    name = name_str.strip()
    if len(name) < 2:
        return False, None, "Имя слишком короткое. Введите хотя бы 2 символа."
    if len(name) > MAX_NAME_LENGTH:
        return False, None, f"Имя слишком длинное (максимум {MAX_NAME_LENGTH} символов)."
    return True, name, ""


def validate_notes(notes_str: str) -> Tuple[bool, str, str]:
    """Validate notes; '-' means "no notes" """
    notes = notes_str.strip()
    if notes == '-':
        return True, '', ""
    if len(notes) > MAX_NOTES_LENGTH:
        return False, '', f"Заметки слишком длинные (максимум {MAX_NOTES_LENGTH} символов)."
    return True, notes, ""


def validate_amount(amount_str: str) -> Tuple[bool, Optional[float], str]:
    """Validate and parse amount string"""
    try:
        # Accept "5 000,50" as well as "5000.50" (regular and non-breaking spaces)
        cleaned = amount_str.replace(',', '.').replace(' ', '').replace(' ', '').strip()
        amount = float(cleaned)
        # float() happily parses 'nan' and 'inf', and nan slips past every comparison
        if not math.isfinite(amount):
            return False, None, "Неверный формат суммы. Используйте числа, например: 5000 или 5000.50"
        if amount <= 0:
            return False, None, "Сумма должна быть больше нуля"
        if amount > 1_000_000_000:
            return False, None, "Сумма слишком большая"
        return True, round(amount, 2), ""
    except ValueError:
        return False, None, "Неверный формат суммы. Используйте числа, например: 5000 или 5000.50"


def validate_interval_value(value_str: str) -> Tuple[bool, Optional[int], str]:
    """Validate interval value"""
    try:
        value = int(value_str.strip())
        if value <= 0:
            return False, None, "Значение должно быть больше нуля"
        if value > 100:
            return False, None, "Значение слишком большое (максимум 100)"
        return True, value, ""
    except ValueError:
        return False, None, "Неверный формат. Введите целое число, например: 1, 2, 3"


def get_interval_description(interval_type: str, interval_value: int) -> str:
    """Get human-readable interval description"""
    interval_map = {
        'monthly': ('месяц', 'месяца', 'месяцев'),
        'weekly': ('неделю', 'недели', 'недель'),
        'yearly': ('год', 'года', 'лет')
    }

    if interval_type not in interval_map:
        return f"{interval_value} {interval_type}"

    words = interval_map[interval_type]

    if interval_value == 1:
        return f"Каждый {words[0]}"
    elif interval_value % 10 == 1 and interval_value % 100 != 11:
        return f"Каждый {interval_value} {words[0]}"
    elif 2 <= interval_value % 10 <= 4 and (interval_value % 100 < 10 or interval_value % 100 >= 20):
        return f"Каждые {interval_value} {words[1]}"
    else:
        return f"Каждые {interval_value} {words[2]}"


async def send_long_message(message: types.Message, text: str, **kwargs):
    """Send long message, splitting if necessary"""
    max_length = 4096

    if len(text) <= max_length:
        await message.answer(text, **kwargs)
    else:
        # Split by lines
        lines = text.split('\n')
        current_message = ""

        for line in lines:
            if len(current_message) + len(line) + 1 <= max_length:
                current_message += line + '\n'
            else:
                if current_message:
                    await message.answer(current_message, **kwargs)
                current_message = line + '\n'

        if current_message:
            await message.answer(current_message, **kwargs)


def generate_excel_report(clients: list, history: list) -> Optional[bytes]:
    """Generate Excel report with clients and payment history.

    Built in memory (no temp files to collide on or clean up). It is CPU-bound
    and synchronous, so call it via asyncio.to_thread from handlers.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        def write_sheet(ws, headers, rows, widths):
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = header_fill
            for row_idx, values in enumerate(rows, 2):
                for col, value in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    # A name or note starting with '=' would otherwise become a live formula
                    if isinstance(value, str):
                        cell.data_type = 's'
            for col, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(col)].width = width

        write_sheet(
            wb.active,
            ["ID", "Имя", "Сумма", "Валюта", "Категория", "Интервал", "След. платеж", "Заметки"],
            ([c['id'], c['name'], c['amount'], c['currency'], c['category'],
              f"{c['interval_value']} {c['interval_type']}", c['next_payment_date'], c['notes']]
             for c in clients),
            [15, 30, 15, 15, 15, 15, 15, 40],
        )
        wb.active.title = "Клиенты"

        write_sheet(
            wb.create_sheet("История платежей"),
            ["Дата", "Клиент", "Сумма", "Валюта", "Статус", "Заметки"],
            ([h['payment_date'], h['client_name'], h['amount'], h['currency'], h['status'], h['notes']]
             for h in history),
            [15, 30, 15, 15, 15, 40],
        )

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except Exception as e:
        logger.error(f"Error generating Excel report: {e}")
        return None
